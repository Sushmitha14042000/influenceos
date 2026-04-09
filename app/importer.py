from __future__ import annotations

from datetime import datetime, timezone
from pathlib import Path

import openpyxl

from app.db import execute, fetch_one
from app.models import Intent, Status

REQUIRED_COLUMNS = [
    "account_id",
    "device_id",
    "post_url",
    "intent",
    "sentiment_tag",
    "comment_template",
    "priority",
    "description",
]


def _utc_now() -> str:
    return datetime.now(timezone.utc).isoformat()


def import_excel(file_path: str) -> int:
    wb = openpyxl.load_workbook(file_path)
    ws = wb.active

    headers = {str(cell.value).strip().lower(): index for index, cell in enumerate(ws[1], start=1)}
    missing = [c for c in REQUIRED_COLUMNS if c not in headers]
    if missing:
        raise ValueError(f"Missing required columns: {', '.join(missing)}")

    execute(
        "INSERT INTO import_batches(source_file, created_at) VALUES (?, ?)",
        [str(Path(file_path).name), _utc_now()],
    )
    batch_row = fetch_one("SELECT id FROM import_batches ORDER BY id DESC LIMIT 1")
    if not batch_row:
        raise RuntimeError("Failed to create import batch")
    batch_id = int(batch_row["id"])

    for row in ws.iter_rows(min_row=2, values_only=True):
        if row is None:
            continue
        account_id = str(row[headers["account_id"] - 1] or "").strip()
        device_id = str(row[headers["device_id"] - 1] or "").strip()
        post_url = str(row[headers["post_url"] - 1] or "").strip()
        intent_raw = str(row[headers["intent"] - 1] or "").strip().lower()
        sentiment_tag = str(row[headers["sentiment_tag"] - 1] or "neutral").strip().lower()
        comment_template = row[headers["comment_template"] - 1]
        description = row[headers["description"] - 1]

        priority_cell = row[headers["priority"] - 1]
        priority = int(priority_cell) if priority_cell is not None else 1

        if not account_id or not post_url or not intent_raw:
            continue

        try:
            intent = Intent(intent_raw)
        except ValueError as exc:
            raise ValueError(f"Invalid intent '{intent_raw}' in file {file_path}") from exc

        execute(
            """
            INSERT INTO jobs(
                batch_id, account_id, device_id, post_url, intent, sentiment_tag,
                comment_template, priority, description, status, created_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            [
                batch_id,
                account_id,
                device_id,
                post_url,
                intent.value,
                sentiment_tag,
                str(comment_template) if comment_template is not None else None,
                priority,
                str(description) if description is not None else None,
                Status.BLOCKED_NO_DEVICE.value if not device_id else Status.PENDING.value,
                _utc_now(),
            ],
        )

    wb.close()
    return batch_id
